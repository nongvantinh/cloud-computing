#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Sinh file trình chiếu 18 slide từ presentation.md.

Nguyên tắc:
  * Slide chỉ giữ Ý CHÍNH, trình bày trực quan (hộp, sơ đồ luồng, bảng so sánh).
  * Toàn bộ diễn giải chi tiết nằm ở SPEAKER NOTES — lấy nguyên văn từ
    presentation.md qua tools/parse_presentation.py.
  * Cấu trúc 18 slide / 6 phần giữ đúng như presentation.md.

Chạy:  .venv/bin/python tools/build_presentation.py
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from parse_presentation import parse  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Chapter11_Vietnamese_Presentation.pptx"

# ----------------------------------------------------------------- palette
NAVY = RGBColor.from_string("0E2D4E")
NAVY2 = RGBColor.from_string("173F68")
BLUE = RGBColor.from_string("1F6FB2")
BLUEL = RGBColor.from_string("5D9BD3")
SKY = RGBColor.from_string("EAF2F9")
SKY2 = RGBColor.from_string("F5F9FC")
BORDER = RGBColor.from_string("C3D8EA")
TEXT = RGBColor.from_string("2B3A47")
MUTED = RGBColor.from_string("64788A")
WHITE = RGBColor.from_string("FFFFFF")
PALE = RGBColor.from_string("BDD3E6")
RED = RGBColor.from_string("B8342A")
REDBG = RGBColor.from_string("FBEDEB")
REDBD = RGBColor.from_string("EEBDB7")
GREEN = RGBColor.from_string("1B7A57")
GRNBG = RGBColor.from_string("E9F5F0")
GRNBD = RGBColor.from_string("B3DCCB")
AMB = RGBColor.from_string("A86A00")
AMBBG = RGBColor.from_string("FDF4E4")
AMBBD = RGBColor.from_string("EBD2A3")

FONT = "Arial"

SW, SH = 13.333, 7.5
ML, MR = 0.62, 0.62
CW = SW - ML - MR
TOP = 1.42                      # mép trên vùng nội dung
BOT = 7.02                      # mép dưới vùng nội dung

warnings = []

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


# =============================================================== helpers
def new_slide():
    return prs.slides.add_slide(BLANK)


def noshadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    noshadow(s)
    if adj is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.text_frame.word_wrap = True
    return s


def txbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def rich(p, text, size, color, bold=False, italic=False, hl=None, hl_bold=True,
         font=FONT, space_after=0, space_before=0, align=PP_ALIGN.LEFT,
         line_spacing=1.16):
    """`**...**` -> phần được nhấn mạnh (đậm + màu hl)."""
    p.alignment = align
    p.line_spacing = line_spacing
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    hl = color if hl is None else hl
    for i, part in enumerate(text.split("**")):
        if part == "":
            continue
        r = p.add_run()
        r.text = part
        f = r.font
        f.name = font
        f.size = Pt(size)
        strong = (i % 2 == 1)
        f.bold = bold or (strong and hl_bold)
        f.italic = italic
        f.color.rgb = hl if strong else color
    return p


def para(tf, first=False):
    return tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()


CHAR_W = 0.50


def n_lines(text, w_in, size):
    plain = text.replace("**", "")
    cpl = max(6, int(w_in * 72.0 / (size * CHAR_W)))
    words, line, cnt = plain.split(), 0, 1
    for wd in words:
        add = len(wd) + (1 if line else 0)
        if line + add > cpl:
            cnt += 1
            line = len(wd)
        else:
            line += add
    return cnt


def est_height(lines, w_in, size, gap, line_spacing=1.16):
    total = 0.0
    for ln in lines:
        total += n_lines(ln, w_in, size) * size * line_spacing / 72.0
    total += gap * max(0, len(lines) - 1) / 72.0
    return total + 0.03


def textblock(slide, x, y, w, lines, size=11, color=TEXT, hl=NAVY, gap=5,
              line_spacing=1.16, align=PP_ALIGN.LEFT, italic=False, bold=False,
              anchor=MSO_ANCHOR.TOP, h=None, font=FONT):
    est = est_height(lines, w, size, gap, line_spacing)
    tf = txbox(slide, x, y, w, h if h else est, anchor=anchor)
    for i, ln in enumerate(lines):
        p = para(tf, first=(i == 0))
        rich(p, ln, size, color, hl=hl, align=align, italic=italic, bold=bold,
             font=font, space_after=gap if i < len(lines) - 1 else 0,
             line_spacing=line_spacing)
    return est


def bullets(slide, x, y, w, items, size=11, color=TEXT, hl=NAVY, gap=6,
            mark="▪", mark_color=BLUE, indent=0.19, line_spacing=1.16):
    cur = y
    for it in items:
        mtf = txbox(slide, x, cur - 0.004, indent, 0.3)
        rich(mtf.paragraphs[0], mark, size, mark_color, bold=True,
             line_spacing=line_spacing)
        h = est_height([it], w - indent, size, 0, line_spacing)
        tf = txbox(slide, x + indent, cur, w - indent, h)
        rich(tf.paragraphs[0], it, size, color, hl=hl, line_spacing=line_spacing)
        cur += h + gap / 72.0
    return cur - y - gap / 72.0


def check(name, y_end, limit=BOT):
    if y_end > limit + 0.01:
        warnings.append(f"{name}: tràn {y_end - limit:.2f}in")


# ============================================================= components
def card_h(w, heading, lines, head_size=12.5, body_size=10.8, bullet=False,
           gap=5, pad=0.15, accent=None):
    """Chiều cao tối thiểu để chứa hết nội dung của một hộp."""
    pw = w - 2 * pad - (0.07 if accent else 0)
    h = 2 * pad
    if heading:
        h += est_height([heading], pw, head_size, 0, 1.1) + 0.07
    if lines:
        if bullet:
            h += sum(est_height([l], pw - 0.15, body_size, 0, 1.16) for l in lines)
            h += gap * (len(lines) - 1) / 72.0
        else:
            h += est_height(lines, pw, body_size, gap, 1.16)
    return h


def card(slide, x, y, w, h, heading, lines, fill=SKY2, border=BORDER,
         head_color=NAVY, body_color=TEXT, hl=BLUE, head_size=12.5,
         body_size=10.8, bullet=False, gap=5, pad=0.15, accent=None, name=""):
    """Hộp nội dung: tiêu đề + các dòng (hoặc gạch đầu dòng).

    `h` là chiều cao mong muốn; hộp tự nới ra nếu nội dung cần nhiều hơn.
    """
    h = max(h or 0.0, card_h(w, heading, lines, head_size, body_size, bullet,
                             gap, pad, accent))
    rect(slide, x, y, w, h, fill=fill, line=border, lw=1.0, adj=0.06)
    if accent:
        rect(slide, x, y, 0.065, h, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    px = x + pad + (0.07 if accent else 0)
    pw = w - 2 * pad - (0.07 if accent else 0)
    cur = y + pad
    if heading:
        hh = est_height([heading], pw, head_size, 0, 1.1)
        tf = txbox(slide, px, cur, pw, hh)
        rich(tf.paragraphs[0], heading, head_size, head_color, bold=True,
             hl=head_color, line_spacing=1.1)
        cur += hh + 0.07
    if lines:
        if bullet:
            bullets(slide, px, cur, pw, lines, size=body_size,
                    color=body_color, hl=hl, gap=gap, mark="•",
                    mark_color=hl, indent=0.15)
        else:
            textblock(slide, px, cur, pw, lines, size=body_size,
                      color=body_color, hl=hl, gap=gap)
    return y + h


def cards_row(slide, x, y, w, h, specs, gap=0.22, **kw):
    """specs: [(heading, [lines], accent|None, {tùy chọn riêng}), ...]

    Mọi hộp trong hàng dùng chung chiều cao = max(h, nội dung cao nhất).
    """
    n = len(specs)
    cw = (w - gap * (n - 1)) / n
    need = h or 0.0
    for sp in specs:
        kwargs = dict(kw)
        if len(sp) > 3 and isinstance(sp[3], dict):
            kwargs.update(sp[3])
        accent = sp[2] if len(sp) > 2 else None
        need = max(need, card_h(
            cw, sp[0], sp[1], kwargs.get("head_size", 12.5),
            kwargs.get("body_size", 10.8), kwargs.get("bullet", False),
            kwargs.get("gap", 5), kwargs.get("pad", 0.15), accent))
    for i, sp in enumerate(specs):
        kwargs = dict(kw)
        if len(sp) > 3 and isinstance(sp[3], dict):
            kwargs.update(sp[3])
        accent = sp[2] if len(sp) > 2 else None
        card(slide, x + i * (cw + gap), y, cw, need, sp[0], sp[1],
             accent=accent, **kwargs)
    return y + need


def chips(slide, x, y, w, items, size=10.8, h=0.34, gap=0.16, fill=SKY,
          border=BORDER, color=NAVY, per_row=None):
    """Dãy “viên” từ khóa. Trả về mép dưới."""
    n = len(items)
    per_row = per_row or n
    rows = [items[i:i + per_row] for i in range(0, n, per_row)]
    cur = y
    for row in rows:
        cwid = (w - gap * (len(row) - 1)) / len(row)
        for i, it in enumerate(row):
            s = rect(slide, x + i * (cwid + gap), cur, cwid, h, fill=fill,
                     line=border, lw=0.75, adj=0.30)
            tf = s.text_frame
            tf.margin_left = tf.margin_right = Inches(0.04)
            tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            rich(tf.paragraphs[0], it, size, color, hl=color,
                 align=PP_ALIGN.CENTER, line_spacing=1.0)
        cur += h + 0.10
    return cur - 0.10


def arrow_right(slide, x, y, w=0.24, h=0.15, color=BLUEL):
    rect(slide, x, y, w, h, fill=color, shape=MSO_SHAPE.RIGHT_ARROW)


def arrow_down(slide, x, y, w=0.15, h=0.20, color=BLUEL):
    rect(slide, x, y, w, h, fill=color, shape=MSO_SHAPE.DOWN_ARROW)


def flow_h(slide, x, y, w, h, items, gap=0.32, size=10.8, fill=SKY,
           border=BLUEL, color=NAVY, last=None, widths=None):
    """Chuỗi hộp nối bằng mũi tên. Mỗi item: 'Tiêu đề|dòng phụ'."""
    n = len(items)
    if widths:
        tot = sum(widths)
        ws = [(w - gap * (n - 1)) * v / tot for v in widths]
    else:
        ws = [(w - gap * (n - 1)) / n] * n
    for i, it in enumerate(items):
        parts = it.split("|")
        need = sum(n_lines(p_, ws[i] - 0.12, size if k == 0 else size - 1.6)
                   * (size if k == 0 else size - 1.6) * 1.06 / 72.0
                   for k, p_ in enumerate(parts)) + 0.12
        h = max(h, need)
    cx = x
    for i, it in enumerate(items):
        f, b, c = fill, border, color
        if i == n - 1 and last:
            f, b, c = last
        s = rect(slide, cx, y, ws[i], h, fill=f, line=b, lw=1.0, adj=0.10)
        tf = s.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        parts = it.split("|")
        for j, ln in enumerate(parts):
            p = para(tf, first=(j == 0))
            rich(p, ln, size if j == 0 else size - 1.6,
                 c if j == 0 else MUTED, bold=(j == 0), hl=c,
                 align=PP_ALIGN.CENTER, line_spacing=1.06,
                 space_before=2 if j else 0)
        if i < n - 1:
            arrow_right(slide, cx + ws[i] + (gap - 0.24) / 2, y + h / 2 - 0.075)
        cx += ws[i] + gap
    return y + h


def flow_v(slide, x, y, w, items, bh=0.36, gap=0.20, size=10.6, fill=SKY,
           border=BLUEL, color=NAVY, last=None):
    """Chuỗi hộp theo chiều dọc, nối bằng mũi tên xuống."""
    cur = y
    n = len(items)
    bh = max(bh, max(n_lines(it, w - 0.12, size) * size * 1.05 / 72.0 + 0.10
                     for it in items))
    ah = max(0.11, min(0.18, gap - 0.03))
    for i, it in enumerate(items):
        f, b, c = fill, border, color
        if i == n - 1 and last:
            f, b, c = last
        s = rect(slide, x, cur, w, bh, fill=f, line=b, lw=1.0, adj=0.14)
        tf = s.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        rich(tf.paragraphs[0], it, size, c, hl=c, align=PP_ALIGN.CENTER,
             line_spacing=1.05)
        cur += bh
        if i < n - 1:
            arrow_down(slide, x + w / 2 - 0.075, cur + (gap - ah) / 2, h=ah)
            cur += gap
    return cur


def steps(slide, x, y, w, h, items, gap=0.22, size=10.8, circ=0.30, start=1):
    """Các bước đánh số nằm ngang: 'Tiêu đề|mô tả'."""
    n = len(items)
    cwid = (w - gap * (n - 1)) / n
    tw = cwid - 0.26 - circ - 0.10
    for it in items:
        parts = it.split("|")
        need = est_height([parts[0]], tw, size, 0, 1.08)
        if len(parts) > 1:
            need += est_height([parts[1]], tw, size - 1.4, 0, 1.1) + 0.04
        h = max(h, need + 0.26)
    for i, it in enumerate(items):
        bx = x + i * (cwid + gap)
        rect(slide, bx, y, cwid, h, fill=SKY2, line=BORDER, lw=1.0, adj=0.08)
        c = rect(slide, bx + 0.13, y + 0.13, circ, circ, fill=BLUE,
                 shape=MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        rich(tf.paragraphs[0], str(start + i), size - 1.5, WHITE, bold=True,
             align=PP_ALIGN.CENTER, line_spacing=1.0)
        parts = it.split("|")
        tf = txbox(slide, bx + 0.13 + circ + 0.10, y + 0.13, tw, h - 0.26)
        rich(tf.paragraphs[0], parts[0], size, NAVY, bold=True, hl=NAVY,
             line_spacing=1.08)
        if len(parts) > 1:
            rich(tf.add_paragraph(), parts[1], size - 1.4, TEXT, hl=BLUE,
                 line_spacing=1.1, space_before=3)
    return y + h


def grid(slide, x, y, w, header, rows, ratios=None, size=10.4, head_size=10.6,
         pad=0.10, first_col_bold=True):
    """Bảng vẽ bằng hình khối — chiều cao và màu sắc kiểm soát được."""
    ncol = len(header) if header else len(rows[0])
    ratios = ratios or [1] * ncol
    tot = sum(ratios)
    ws = [w * r / tot for r in ratios]
    xs, acc = [], x
    for wd in ws:
        xs.append(acc)
        acc += wd
    cur = y
    if header:
        hh = max(est_height([c], ws[i] - 2 * pad, head_size, 0, 1.1)
                 for i, c in enumerate(header)) + 0.13
        rect(slide, x, cur, w, hh, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
        for i, c in enumerate(header):
            tf = txbox(slide, xs[i] + pad, cur, ws[i] - 2 * pad, hh,
                       anchor=MSO_ANCHOR.MIDDLE)
            rich(tf.paragraphs[0], c, head_size, WHITE, bold=True, hl=WHITE,
                 line_spacing=1.1)
        cur += hh
    for r, row in enumerate(rows):
        rh = max(est_height([c], ws[i] - 2 * pad, size, 0, 1.14)
                 for i, c in enumerate(row)) + 0.12
        rect(slide, x, cur, w, rh, fill=WHITE if r % 2 == 0 else SKY2,
             line=BORDER, lw=0.5, shape=MSO_SHAPE.RECTANGLE)
        for i, c in enumerate(row):
            tf = txbox(slide, xs[i] + pad, cur, ws[i] - 2 * pad, rh,
                       anchor=MSO_ANCHOR.MIDDLE)
            bold = first_col_bold and i == 0
            rich(tf.paragraphs[0], c, size, NAVY if bold else TEXT, bold=bold,
                 hl=BLUE if not bold else NAVY, line_spacing=1.14)
        cur += rh
    return cur


def callout(slide, x, y, w, text, tone="blue", size=11.4, pad=0.13):
    """Hộp thông điệp nhấn mạnh; trả về mép dưới."""
    palette = {"blue": (SKY, BORDER, NAVY, BLUE),
               "red": (REDBG, REDBD, RED, RED),
               "green": (GRNBG, GRNBD, GREEN, GREEN),
               "amber": (AMBBG, AMBBD, AMB, AMB),
               "navy": (NAVY, NAVY, WHITE, BLUEL)}
    fill, bd, color, accent = palette[tone]
    inner = w - 2 * pad - 0.10
    h = est_height([text], inner, size, 0, 1.18) + 2 * pad
    rect(slide, x, y, w, h, fill=fill, line=bd, lw=1.0, adj=0.10)
    rect(slide, x, y, 0.07, h, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    tf = txbox(slide, x + pad + 0.07, y + pad, inner, h - 2 * pad)
    rich(tf.paragraphs[0], text, size, color, hl=color, line_spacing=1.18)
    return y + h


def layers(slide, x, y, w, items, bh=0.40, gap=0.09, size=11, shrink=0.0,
           fill=SKY, border=BLUEL, color=NAVY, top_tone=None):
    """Các lớp xếp chồng (kiến trúc / defense-in-depth)."""
    cur = y
    n = len(items)
    for i, it in enumerate(items):
        inset = shrink * i * w / max(1, n)
        f, b, c = fill, border, color
        if top_tone and i == n - 1:
            f, b, c = top_tone
        s = rect(slide, x + inset / 2, cur, w - inset, bh, fill=f, line=b,
                 lw=1.0, adj=0.12)
        tf = s.text_frame
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        parts = it.split("|")
        rich(tf.paragraphs[0], parts[0], size, c, bold=True, hl=c,
             align=PP_ALIGN.CENTER, line_spacing=1.04)
        if len(parts) > 1:
            rich(tf.add_paragraph(), parts[1], size - 1.8, MUTED,
                 align=PP_ALIGN.CENTER, line_spacing=1.04, space_before=1)
        cur += bh + gap
    return cur - gap


def label(slide, x, y, w, text, size=11.5, color=NAVY, bar=True):
    """Nhãn mục nhỏ kèm vạch màu; trả về mép dưới."""
    h = est_height([text], w - 0.16, size, 0, 1.1)
    if bar:
        rect(slide, x, y + 0.02, 0.055, max(0.15, h - 0.05), fill=BLUE,
             shape=MSO_SHAPE.RECTANGLE)
    tf = txbox(slide, x + (0.13 if bar else 0), y, w - 0.13, h)
    rich(tf.paragraphs[0], text, size, color, bold=True, hl=color,
         line_spacing=1.1)
    return y + h + 0.07


# ================================================================== khung
KICKER = {
    2: "NỘI DUNG TRÌNH BÀY",
    3: "PHẦN 1 · NỀN TẢNG BẢO MẬT ĐÁM MÂY",
    4: "PHẦN 2 · BẢO VỆ DỮ LIỆU TRÊN CLOUD",
    5: "PHẦN 3 · MẬT MÃ NÂNG CAO CHO CLOUD",
    6: "PHẦN 3 · MẬT MÃ NÂNG CAO CHO CLOUD",
    7: "PHẦN 3 · MẬT MÃ NÂNG CAO CHO CLOUD",
    8: "PHẦN 3 · MẬT MÃ NÂNG CAO CHO CLOUD",
    9: "PHẦN 3 · MẬT MÃ NÂNG CAO CHO CLOUD",
    10: "PHẦN 3 · MẬT MÃ NÂNG CAO CHO CLOUD",
    11: "PHẦN 4 · BẢO VỆ HẠ TẦNG & VÙNG TIN CẬY",
    12: "PHẦN 5 · BẢO MẬT ỨNG DỤNG & MÁY CHỦ",
    13: "PHẦN 5 · BẢO MẬT ỨNG DỤNG & MÁY CHỦ",
    14: "PHẦN 5 · BẢO MẬT ỨNG DỤNG & MÁY CHỦ",
    15: "PHẦN 5 · BẢO MẬT ỨNG DỤNG & MÁY CHỦ",
    16: "PHẦN 5 · BẢO MẬT ỨNG DỤNG & MÁY CHỦ",
    17: "PHẦN 6 · KIỂM SOÁT & PHÒNG THỦ ĐA LỚP",
    18: "TỔNG KẾT",
}

TITLES = {
    2: "MỤC LỤC TRÌNH BÀY",
    3: "NỀN TẢNG BẢO MẬT ĐÁM MÂY",
    4: "BẢO VỆ DỮ LIỆU TRÊN CLOUD",
    5: "MẬT MÃ NÂNG CAO CHO CLOUD",
    6: "ABE — MÃ HÓA DỰA TRÊN THUỘC TÍNH",
    7: "KP-ABE PHI TẬP TRUNG & CHỨNG MINH BẢO MẬT",
    8: "FHE — TÍNH TOÁN TRÊN DỮ LIỆU MÃ HÓA",
    9: "FHE — TÍNH TOÁN THUÊ NGOÀI & ỨNG DỤNG Y TẾ",
    10: "SEARCHABLE ENCRYPTION — TÌM KIẾM TRÊN DỮ LIỆU MÃ HÓA",
    11: "BẢO VỆ HẠ TẦNG & VÙNG TIN CẬY",
    12: "IaaS / PaaS / SaaS — RANH GIỚI TRÁCH NHIỆM",
    13: "BẢO MẬT ỨNG DỤNG PaaS",
    14: "ĐỊNH DANH & QUẢN LÝ TRUY CẬP (IAM & SSO)",
    15: "BẢO MẬT SaaS & KIỂM SOÁT TRUY CẬP CHI TIẾT",
    16: "BẢO MẬT MÁY CHỦ ẢO (VIRTUAL MACHINE SECURITY)",
    17: "SECURITY CONTROLS — NGĂN NGỪA, PHÁT HIỆN & KHẮC PHỤC",
    18: "TỔNG KẾT — KIẾN TRÚC BẢO MẬT & DEFENSE-IN-DEPTH",
}


def head(slide, kicker, title):
    rect(slide, ML, 0.42, 0.085, 0.66, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    if kicker:
        tf = txbox(slide, ML + 0.22, 0.36, CW - 0.3, 0.24)
        rich(tf.paragraphs[0], kicker, 10.5, BLUE, bold=True, line_spacing=1.0)
    size = 24 if len(title) <= 44 else (21 if len(title) <= 56 else 18.5)
    tf = txbox(slide, ML + 0.22, 0.62, CW - 0.3, 0.56)
    rich(tf.paragraphs[0], title, size, NAVY, bold=True, line_spacing=1.02)
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ML),
                                    Inches(1.26), Inches(SW - MR), Inches(1.26))
    ln.line.color.rgb = BORDER
    ln.line.width = Pt(1.0)


def footer(slide, idx, total=18):
    tf = txbox(slide, ML, 7.13, 8.0, 0.24)
    rich(tf.paragraphs[0], "Bảo mật đám mây (Cloud Security) · Nông Văn Tình",
         9, MUTED, line_spacing=1.0)
    tf = txbox(slide, SW - MR - 1.6, 7.13, 1.6, 0.24)
    rich(tf.paragraphs[0], f"{idx}/{total}", 9, MUTED, align=PP_ALIGN.RIGHT,
         line_spacing=1.0)


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text.strip()


def frame(d):
    s = new_slide()
    head(s, KICKER.get(d["num"], ""), TITLES.get(d["num"], d["title"]))
    footer(s, d["num"])
    notes(s, d["notes"])
    return s


# ================================================================ SLIDE 1
def slide1(d):
    s = new_slide()
    rect(s, 0, 0, SW, SH, fill=NAVY, shape=MSO_SHAPE.RECTANGLE)
    rect(s, SW - 4.3, -1.7, 5.6, 5.6, fill=NAVY2, shape=MSO_SHAPE.OVAL)
    rect(s, SW - 2.6, 4.2, 4.4, 4.4, fill=NAVY2, shape=MSO_SHAPE.OVAL)
    rect(s, 0, 0, 0.14, SH, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    tf = txbox(s, 1.15, 1.78, 9.0, 0.4)
    rich(tf.paragraphs[0], "CHƯƠNG 11", 17, BLUEL, bold=True, line_spacing=1.0)
    tf = txbox(s, 1.15, 2.26, 10.4, 0.9)
    rich(tf.paragraphs[0], "BẢO MẬT ĐÁM MÂY", 48, WHITE, bold=True,
         hl=WHITE, line_spacing=1.0)
    tf = txbox(s, 1.15, 3.36, 10.0, 0.45)
    rich(tf.paragraphs[0], "Cloud Security", 20, BLUEL, italic=True,
         line_spacing=1.0)
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.15),
                                Inches(4.06), Inches(3.4), Inches(4.06))
    ln.line.color.rgb = BLUE
    ln.line.width = Pt(3)
    tf = txbox(s, 1.15, 4.34, 10.2, 0.4)
    rich(tf.paragraphs[0],
         "Dữ liệu · Mã hóa · Hạ tầng · PaaS · SaaS · Máy chủ ảo · Kiểm soát bảo mật",
         13.5, PALE, line_spacing=1.2)
    tf = txbox(s, 1.15, 4.92, 10.2, 0.5)
    rich(tf.paragraphs[0],
         "“Khi đưa hệ thống lên Cloud, mô hình bảo mật cũng hoàn toàn thay đổi.”",
         13, BLUEL, italic=True, line_spacing=1.2)
    tf = txbox(s, 1.15, 6.18, 8.0, 0.35)
    rich(tf.paragraphs[0], "Người thực hiện: **Nông Văn Tình**", 13, PALE,
         hl=WHITE, line_spacing=1.1)
    tf = txbox(s, SW - MR - 1.6, 7.13, 1.6, 0.24)
    rich(tf.paragraphs[0], "1/18", 9, RGBColor.from_string("6E8FAE"),
         align=PP_ALIGN.RIGHT, line_spacing=1.0)
    notes(s, d["notes"])


# ================================================================ SLIDE 2
def slide2(d):
    s = frame(d)
    specs = [
        ("1 · NỀN TẢNG BẢO MẬT ĐÁM MÂY",
         ["Cloud Security & vấn đề đặc thù", "Shared Responsibility Model",
          "Bộ ba bảo mật CIA"], BLUE),
        ("2 · BẢO VỆ DỮ LIỆU TRÊN CLOUD",
         ["Vòng đời & trạng thái dữ liệu", "Data-in-Transit & Metadata",
          "Privacy · Integrity · Location"], BLUE),
        ("3 · MẬT MÃ NÂNG CAO CHO CLOUD",
         ["Attribute-Based Encryption (ABE)",
          "Fully Homomorphic Encryption (FHE)",
          "Searchable Encryption (SE)"], BLUE),
        ("4 · BẢO VỆ HẠ TẦNG & VÙNG TIN CẬY",
         ["Public Cloud vs Private Cloud", "Network Security & cấu hình sai",
          "Trust Zones & IAM"], BLUEL),
        ("5 · BẢO MẬT ỨNG DỤNG & MÁY CHỦ",
         ["IaaS · PaaS · SaaS", "Access Control · SSO · Hardening",
          "Virtual Server Security"], BLUEL),
        ("6 · KIỂM SOÁT & PHÒNG THỦ ĐA LỚP",
         ["Deterrent · Preventive", "Detective · Corrective",
          "Defense-in-Depth"], BLUEL),
    ]
    h = 1.66
    cards_row(s, ML, TOP + 0.24, CW, h, specs[:3], gap=0.26, bullet=True,
              head_size=12.5, body_size=11)
    cards_row(s, ML, TOP + 0.24 + h + 0.30, CW, h, specs[3:], gap=0.26,
              bullet=True, head_size=12.5, body_size=11)
    y = callout(s, ML, TOP + 0.24 + 2 * h + 0.62, CW,
                "Mạch bài: từ **trách nhiệm và mục tiêu bảo mật** → **bảo vệ dữ liệu** "
                "→ **mật mã nâng cao** → **hạ tầng, ứng dụng, máy chủ** → "
                "**các lớp kiểm soát**.")
    check("Slide 2", y)


# ================================================================ SLIDE 3
def slide3(d):
    s = frame(d)
    wl = 6.00
    y = TOP
    y = label(s, ML, y, wl, "1 · Cloud Security là gì?")
    y = card(s, ML, y, wl, 0.82, None,
             ["Tập hợp **công nghệ · chính sách · quy trình · biện pháp kiểm soát** "
              "nhằm bảo vệ tài nguyên trên Cloud."],
             body_size=11.0, pad=0.14)
    y = chips(s, ML, y + 0.14, wl,
              ["Dữ liệu", "Ứng dụng", "Hạ tầng", "Danh tính"], size=10.8)
    y += 0.20
    y = label(s, ML, y, wl, "2 · Shared Responsibility Model")
    y = grid(s, ML, y, wl,
             ["CSP — Security OF the Cloud", "Khách hàng — Security IN the Cloud"],
             [["Trung tâm dữ liệu, phần cứng, mạng nền tảng", "Cấu hình tài nguyên"],
              ["Hạ tầng / dịch vụ do CSP cung cấp", "Danh tính & quyền truy cập"],
              ["Bảo mật nền tảng Cloud", "Dữ liệu, ứng dụng & biện pháp tự triển khai"]],
             first_col_bold=False, size=10.4)
    y = callout(s, ML, y + 0.14, wl,
                "Ranh giới trách nhiệm **thay đổi theo IaaS → PaaS → SaaS**.", size=11)
    y = callout(s, ML, y + 0.16, wl,
                "**Dữ liệu bị rò rỉ — ai chịu trách nhiệm?** Không có câu trả lời duy "
                "nhất: cần xác định **nguyên nhân sự cố** và **ranh giới trách nhiệm** "
                "của dịch vụ đang dùng.", tone="red", size=11)

    x2 = ML + wl + 0.32
    w2 = CW - wl - 0.32
    yb = TOP
    yb = label(s, x2, yb, w2, "3 · Mục tiêu bảo mật: CIA Triad")
    yb = cards_row(s, x2, yb, w2, 1.58, [
        ("Confidentiality", ["Chỉ người được phép mới truy cập được dữ liệu."], BLUE),
        ("Integrity", ["Dữ liệu không bị thay đổi trái phép."], BLUE),
        ("Availability", ["Sẵn sàng khi người dùng hợp lệ cần dùng."], BLUE),
    ], gap=0.18, head_size=11.6, body_size=10.6)
    yb = callout(s, x2, yb + 0.18, w2,
                 "**Thách thức trong Cloud:** bên thứ ba + truy cập qua mạng + tài nguyên "
                 "dùng chung + cấu hình động → duy trì CIA khó hơn.",
                 tone="amber", size=11)
    yb = callout(s, x2, yb + 0.20, w2,
                 "Cloud Security = **bảo vệ tài nguyên Cloud** + **xác định đúng trách "
                 "nhiệm** + **duy trì CIA**.", tone="navy", size=11.5)
    check("Slide 3", max(y, yb))


# ================================================================ SLIDE 4
def slide4(d):
    s = frame(d)
    y = TOP
    y = label(s, ML, y, CW, "1 · Dữ liệu trên Cloud đi qua đâu?")
    y = flow_h(s, ML, y, CW, 0.44,
               ["User", "Network", "Cloud Storage", "Cloud Processing", "User"],
               gap=0.30, size=11.4)
    y += 0.16
    y = cards_row(s, ML, y, CW, 1.22, [
        ("Data-in-Transit", ["Đang truyền qua mạng — Internet / mạng riêng",
                             "**Nguy cơ:** nghe lén, can thiệp"], BLUE),
        ("Data-at-Rest", ["Đang lưu trữ — database / cloud storage",
                          "**Nguy cơ:** truy cập trái phép, mất dữ liệu"], BLUE),
        ("Data-in-Use", ["Đang được xử lý — RAM / CPU",
                         "**Nguy cơ:** lộ dữ liệu khi xử lý"], BLUE),
    ], gap=0.24, head_size=12.2, body_size=10.6)
    y += 0.20

    w2 = (CW - 0.30) / 2
    ya = label(s, ML, y, w2, "2 · Bảo vệ dữ liệu trên đường truyền")
    ya = card(s, ML, ya, w2, 1.16, None,
              ["**Nguy cơ:** nghe lén · can thiệp · giả mạo",
               "**Nguyên tắc:** dùng giao thức và thuật toán mã hóa "
               "**đã được kiểm chứng**"],
              body_size=10.8, pad=0.14)
    ya = callout(s, ML, ya + 0.14, w2,
                 "**Private network ≠ mặc nhiên tin cậy** — dữ liệu cần được bảo vệ dù đi "
                 "qua Internet hay mạng riêng.", tone="amber", size=10.8)

    xb = ML + w2 + 0.30
    yb = label(s, xb, y, w2, "3 · Metadata cũng có giá trị")
    yb = card(s, xb, yb, w2, 1.16, None,
              ["Metadata = thông tin **xung quanh** dữ liệu: ai truy cập, khi nào, "
               "tần suất, dữ liệu nằm ở đâu, khối lượng bao nhiêu.",
               "Cần biết CSP **thu thập gì — bảo vệ ra sao — khách hàng kiểm soát đến đâu**."],
              body_size=10.8, pad=0.14)
    yb = label(s, xb, yb + 0.14, w2, "4 · Bốn vấn đề cần kiểm soát")
    yb = chips(s, xb, yb, w2,
               ["Integrity", "Privacy", "Data Location", "Availability"],
               size=10.2, per_row=4, gap=0.12)

    y = max(ya, yb) + 0.18
    y = callout(s, ML, y, CW,
                "**Bài toán tiếp theo:** Cloud cần lưu trữ và tính toán trên dữ liệu, "
                "nhưng ta không muốn Cloud nhìn thấy dữ liệu gốc → **Advanced Cryptography: "
                "ABE · FHE · Searchable Encryption**.", tone="navy", size=11.4)
    check("Slide 4", y)


# ================================================================ SLIDE 5
def slide5(d):
    s = frame(d)
    y = TOP
    y = label(s, ML, y, CW, "1 · Mã hóa truyền thống chưa đủ")
    y = flow_h(s, ML, y, CW, 0.44,
               ["Dữ liệu gốc", "Mã hóa", "Cloud", "Giải mã", "Xử lý"],
               gap=0.28, size=11.2, last=(REDBG, REDBD, RED))
    y = callout(s, ML, y + 0.10, CW,
                "Khi cần xử lý, **bản rõ thường phải xuất hiện** → làm sao vẫn dùng được "
                "khả năng lưu trữ, tính toán và tìm kiếm của Cloud mà **hạn chế việc Cloud "
                "nhìn thấy dữ liệu gốc**?", tone="red", size=11.2)
    y += 0.24
    y = label(s, ML, y, CW, "2 · Ba kỹ thuật — ba bài toán khác nhau")
    y = cards_row(s, ML, y, CW, 2.10, [
        ("ABE — Access", [
            "**Bài toán:** ai được phép truy cập?",
            "**Ý tưởng:** kiểm soát quyền theo **thuộc tính** thay vì danh tính.",
            "**CP-ABE:** policy gắn ciphertext · **KP-ABE:** policy gắn khóa."], BLUE),
        ("FHE — Compute", [
            "**Bài toán:** Cloud tính toán mà không thấy dữ liệu?",
            "**Ý tưởng:** tính trực tiếp trên **ciphertext**.",
            "Người dùng giữ **quyền giải mã** kết quả."], BLUE),
        ("SE — Search", [
            "**Bài toán:** tìm kiếm mà không giải mã toàn bộ?",
            "**Ý tưởng:** truy vấn qua **chỉ mục từ khóa / trapdoor**.",
            "Ví dụ: tìm email chứa **“Urgent”** trong kho đã mã hóa."], BLUE),
    ], gap=0.26, head_size=13.5, body_size=10.8)
    y = callout(s, ML, y + 0.22, CW,
                "**ABE → Access · FHE → Compute · SE → Search** — ba nhu cầu khác nhau "
                "trong cùng bài toán bảo vệ dữ liệu trên Cloud.", tone="navy", size=11.5)
    check("Slide 5", y)


# ================================================================ SLIDE 6
def slide6(d):
    s = frame(d)
    w2 = (CW - 0.34) / 2
    y = TOP
    y = label(s, ML, y, w2, "1 · ABE giải quyết bài toán gì?")
    y = card(s, ML, y, w2, 0.84, None,
             ["Cấp quyền theo **thuộc tính** thay vì danh tính cá nhân — chính sách được "
              "viết trên các thuộc tính của người dùng."],
             body_size=10.9, pad=0.14)
    y = chips(s, ML, y + 0.16, w2, ["Doctor", "Cardiology", "Hospital A"], size=10.8)
    y = flow_h(s, ML, y + 0.16, w2, 0.44,
               ["Access Policy|Doctor AND Cardiology AND Hospital A"], size=10.8)
    y = cards_row(s, ML, y + 0.16, w2, 0.74, [
        ("Thỏa mãn policy", ["→ Giải mã **thành công**"], None,
         dict(fill=GRNBG, border=GRNBD, head_color=GREEN, hl=GREEN)),
        ("Không thỏa mãn", ["→ Giải mã **thất bại**"], None,
         dict(fill=REDBG, border=REDBD, head_color=RED, hl=RED)),
    ], gap=0.20, head_size=11.4, body_size=10.6)
    y = callout(s, ML, y + 0.16, w2,
                "Giải mã thành công khi **tập thuộc tính của người dùng thỏa mãn Access "
                "Policy** — không nhất thiết khớp hoàn toàn.", size=10.9)

    x2 = ML + w2 + 0.34
    yb = TOP
    yb = label(s, x2, yb, w2, "2 · CP-ABE vs. KP-ABE")
    yb = grid(s, x2, yb, w2, ["", "CP-ABE", "KP-ABE"],
              [["Policy nằm ở", "Ciphertext", "User Secret Key"],
               ["Attributes nằm ở", "User Secret Key", "Ciphertext"],
               ["Ai quyết định policy", "Bên mã hóa dữ liệu", "Cơ quan cấp khóa"],
               ["Câu hỏi trực quan", "“Dữ liệu này cho ai?”",
                "“Người này xem được loại dữ liệu nào?”"]],
              ratios=[1.0, 1.1, 1.35], size=10.2)
    yb = callout(s, x2, yb + 0.14, w2,
                 "**CP-ABE: Policy → Ciphertext**     **KP-ABE: Policy → Key**", size=11)
    yb += 0.20
    yb = label(s, x2, yb, w2, "3 · Cơ chế ABE — bốn bước")
    yb = grid(s, x2, yb, w2, None,
              [["① Setup", "Tạo tham số hệ thống & khóa công khai"],
               ["② Key Generation", "Cấp secret key gắn thuộc tính / policy"],
               ["③ Encryption", "Mã hóa dữ liệu kèm policy / thuộc tính"],
               ["④ Decryption", "Đối chiếu thuộc tính của người dùng với policy"]],
              ratios=[0.8, 1.6], size=10.4)
    y = max(y, yb)
    y = callout(s, ML, y + 0.10, CW,
                "ABE biến quyền truy cập thành **bài toán về thuộc tính và chính sách**: "
                "**CP-ABE** — người mã hóa quyết định ai đọc được; **KP-ABE** — chính sách "
                "trong khóa quyết định đọc được loại dữ liệu nào.", tone="navy", size=11.2)
    check("Slide 6", y)


# ================================================================ SLIDE 7
def slide7(d):
    s = frame(d)
    w2 = (CW - 0.34) / 2
    y = TOP
    y = label(s, ML, y, w2, "1 · Từ KP-ABE cơ bản → phi tập trung")
    y = card(s, ML, y, w2, 1.04, None,
             ["**Bài toán:** nhiều **Attribute Authority (AA)** thuộc các tổ chức khác "
              "nhau — quản lý khóa thế nào mà **không phụ thuộc một cơ quan trung tâm**?"],
             body_size=10.9, pad=0.14)
    y = cards_row(s, ML, y + 0.16, w2, 0.92, [
        ("AA₁", ["Thuộc tính tổ chức 1"], None),
        ("AA₂", ["Thuộc tính tổ chức 2"], None),
        ("AA₃", ["Thuộc tính tổ chức 3"], None),
    ], gap=0.16, head_size=11.4, body_size=10.0)
    y = flow_h(s, ML, y + 0.16, w2, 0.44,
               ["Người dùng nhận thành phần khóa từ nhiều AA"], size=10.8)
    y = card(s, ML, y + 0.16, w2, 0.98, None,
             ["Mỗi AA quản lý một tập thuộc tính riêng.",
              "AA có thể **tham gia / rời hệ thống** mà không phải thiết lập lại toàn bộ."],
             body_size=10.6, bullet=True, pad=0.14)
    y = callout(s, ML, y + 0.14, w2,
                "Không còn **một Authority duy nhất** quyết định toàn bộ quyền truy cập.",
                size=10.9)

    x2 = ML + w2 + 0.34
    yb = TOP
    yb = label(s, x2, yb, w2, "2 · Năm thành phần của lược đồ")
    yb = grid(s, x2, yb, w2, None,
              [["① Global Setup", "Tạo tham số hệ thống dùng chung"],
               ["② Authority Setup", "Mỗi AA tạo khóa cho thuộc tính mình quản lý"],
               ["③ Key Issuing", "Người dùng nhận thành phần khóa từ các AA"],
               ["④ Encryption", "Mã hóa dữ liệu kèm tập thuộc tính mô tả"],
               ["⑤ Decryption", "Ghép các thành phần khóa; thỏa policy → giải mã"]],
              ratios=[0.75, 1.55], size=10.3)
    yb += 0.20
    yb = label(s, x2, yb, w2, "3 · Security Game — chứng minh an toàn")
    yb = flow_h(s, x2, yb, w2, 0.78,
                ["① Key Queries|Xin khóa hợp lệ",
                 "② Challenge|Chọn m₀, m₁ → mã hóa một trong hai",
                 "③ More Queries|Không vi phạm policy",
                 "④ Guess|Đoán b′ = 0 hay 1"],
                gap=0.14, size=10.0)
    yb = callout(s, x2, yb + 0.16, w2,
                 "Nếu Adversary **không đạt lợi thế đáng kể**: Pr[b′ = b] ≈ 1/2 → lược đồ "
                 "đạt **Selective-ID Security** trong mô hình đang xét.", size=10.9)
    y = max(y, yb)
    y = callout(s, ML, y + 0.10, CW,
                "KP-ABE phi tập trung giải quyết bài toán **nhiều Authority**; Security "
                "Game cho ta cách **hình thức hóa** để kiểm chứng độ an toàn.",
                tone="navy", size=11.2)
    check("Slide 7", y)


# ================================================================ SLIDE 8
def slide8(d):
    s = frame(d)
    w2 = (CW - 0.34) / 2
    y = TOP
    y = label(s, ML, y, CW, "1 · FHE giải quyết bài toán gì?")
    textblock(s, ML + 0.02, y, w2, ["**Mã hóa thông thường**"], size=10.8,
              color=RED, hl=RED)
    textblock(s, ML + w2 + 0.36, y, w2, ["**Với FHE**"], size=10.8,
              color=GREEN, hl=GREEN)
    y += 0.26
    flow_h(s, ML, y, w2, 0.42, ["Encrypt", "Cloud", "Decrypt", "Compute"],
           gap=0.18, size=10.2, fill=REDBG, border=REDBD, color=RED)
    y = flow_h(s, ML + w2 + 0.34, y, w2, 0.42,
               ["Encrypt", "Compute trên ciphertext", "Decrypt"],
               gap=0.18, size=10.2, fill=GRNBG, border=GRNBD, color=GREEN)
    y = callout(s, ML, y + 0.16, CW,
                "**Nguyên lý cốt lõi:   Dec( Eval( E(m), f ) ) = f(m)**   — tính toán trên "
                "bản mã cho kết quả tương ứng với tính toán trên bản rõ.", size=11.0)
    y += 0.22

    ya = label(s, ML, y, w2, "2 · Kiến trúc triển khai")
    ya = grid(s, ML, ya, w2, ["Thành phần", "Vai trò"],
              [["User", "Sở hữu dữ liệu — mã hóa và giải mã kết quả"],
               ["HC Node", "Thực hiện phép toán đồng cấu trên ciphertext"],
               ["Bootstrapping Node", "Làm mới ciphertext khi nhiễu tăng cao"]],
              ratios=[0.8, 1.5], size=10.1)
    ya = callout(s, ML, ya + 0.14, w2,
                 "Secure Enclave / Remote Attestation thuộc **kiến trúc triển khai cụ "
                 "thể**, không bắt buộc với mọi hệ thống FHE.", tone="amber", size=10.4)

    x2 = ML + w2 + 0.34
    yb = label(s, x2, y, w2, "3 · Luồng xử lý & bài toán nhiễu")
    yb = flow_v(s, x2 + 0.35, yb, w2 - 0.70, [
        "User: mã hóa dữ liệu",
        "HC Node: tính trên ciphertext (noise ↑)",
        "Bootstrapping: làm mới ciphertext",
        "HC Node: tiếp tục tính toán",
        "User: giải mã → kết quả",
    ], bh=0.32, gap=0.13, size=10.1, last=(GRNBG, GRNBD, GREEN))
    yb = callout(s, x2, yb + 0.12, w2,
                 "Mỗi phép toán làm **nhiễu tăng**; vượt ngưỡng → không giải mã đúng. "
                 "**Bootstrapping = làm mới ciphertext để tính tiếp.**", size=10.4)
    y = max(ya, yb)
    y = callout(s, ML, y + 0.16, CW,
                "**ABE:** ai được truy cập?     **FHE:** Cloud tính toán được gì mà không "
                "nhìn thấy dữ liệu?", tone="navy", size=11.2)
    check("Slide 8", y)


# ================================================================ SLIDE 9
def slide9(d):
    s = frame(d)
    y = TOP
    y = label(s, ML, y, CW, "1 · Tính toán thuê ngoài — ví dụ 5 + 10")
    y = flow_h(s, ML, y, CW, 0.80, [
        "Doanh nghiệp B|Dữ liệu gốc: 5 , 10",
        "① Encrypt|Enc(5), Enc(10)",
        "② Cloud|Tính trên ciphertext",
        "Enc(5 + 10)|Kết quả vẫn mã hóa",
        "③ Decrypt|Kết quả: 15",
    ], gap=0.26, size=11.0, last=(GRNBG, GRNBD, GREEN))
    y += 0.18
    y = cards_row(s, ML, y, CW, 0.94, [
        ("Cloud biết", ["Ciphertext + phép tính cần thực hiện"], None),
        ("Cloud không cần biết", ["Dữ liệu gốc 5 , 10"], None,
         dict(fill=REDBG, border=REDBD, head_color=RED, hl=RED)),
        ("User biết", ["Dữ liệu gốc + kết quả sau giải mã"], None,
         dict(fill=GRNBG, border=GRNBD, head_color=GREEN, hl=GREEN)),
    ], gap=0.24, head_size=11.8, body_size=10.6)
    y += 0.22

    w2 = (CW - 0.34) / 2
    ya = label(s, ML, y, w2, "2 · Ứng dụng: phân tích dữ liệu y tế")
    ya = card(s, ML, ya, w2, 1.06, None,
              ["**Mâu thuẫn:** dữ liệu càng chi tiết → phân tích càng hữu ích, nhưng càng "
               "nhạy cảm → yêu cầu bảo vệ càng cao."],
              fill=AMBBG, border=AMBBD, hl=AMB, body_size=10.9, pad=0.14)
    ya = callout(s, ML, ya + 0.14, w2,
                 "Cloud cung cấp **năng lực tính toán**, dữ liệu bệnh nhân **vẫn ở dạng mã "
                 "hóa** trong suốt quá trình xử lý.", tone="green", size=10.9)

    x2 = ML + w2 + 0.34
    yb = label(s, x2, y, w2, "Luồng xử lý với FHE")
    yb = flow_h(s, x2, yb, w2, 0.94, [
        "Bệnh viện|Mã hóa dữ liệu y tế",
        "Cloud|Phân tích trên ciphertext",
        "Bệnh viện|Giải mã kết quả",
    ], gap=0.20, size=10.6, last=(GRNBG, GRNBD, GREEN))
    y = max(ya, yb)
    y = callout(s, ML, y + 0.20, CW,
                "**FHE tách “quyền tính toán” khỏi “quyền nhìn thấy dữ liệu”.** Vậy Cloud "
                "**tìm kiếm** thế nào khi không đọc được nội dung? → **Searchable "
                "Encryption**.", tone="navy", size=11.2)
    check("Slide 9", y)


# =============================================================== SLIDE 10
def slide10(d):
    s = frame(d)
    y = TOP
    y = label(s, ML, y, CW, "1 · Bài toán: mã hóa rồi thì tìm kiếm thế nào?")
    y = callout(s, ML, y, CW,
                "Mã hóa thông thường bảo vệ nội dung nhưng khiến Cloud **không thể tìm kiếm "
                "trực tiếp**. **Searchable Encryption** cho phép truy vấn theo **từ khóa** "
                "mà không cần công khai nội dung gốc.", size=11.2)
    y += 0.22
    y = label(s, ML, y, CW, "2 · Cơ chế: bốn bước")
    y = steps(s, ML, y, CW, 1.06, [
        "Mã hóa|Chủ sở hữu mã hóa tài liệu & xây chỉ mục từ khóa",
        "Tạo truy vấn|Người dùng sinh **trapdoor** cho từ khóa",
        "Đối sánh|Cloud so khớp trapdoor với chỉ mục",
        "Trả kết quả|Trả tài liệu phù hợp để giải mã theo quyền",
    ], gap=0.20, size=10.4)
    y += 0.22

    w2 = (CW - 0.34) / 2
    ya = label(s, ML, y, w2, "3 · Ví dụ: tìm email chứa “Urgent”")
    ya = flow_v(s, ML + 0.45, ya, w2 - 0.90, [
        "Email đã mã hóa + chỉ mục từ khóa",
        "Trapdoor: “Urgent”",
        "Cloud đối sánh chỉ mục",
        "Trả về email phù hợp",
    ], bh=0.32, gap=0.16, size=10.2, last=(GRNBG, GRNBD, GREEN))

    x2 = ML + w2 + 0.34
    yb = label(s, x2, y, w2, "Hai hướng triển khai phổ biến")
    yb = cards_row(s, x2, yb, w2, 1.10, [
        ("Symmetric-key SE", ["Khóa đối xứng + chỉ mục từ khóa",
                              "Hướng tới **hiệu năng tìm kiếm cao**"], None),
        ("PEKS", ["Public Key Encryption with Keyword Search",
                  "Tìm kiếm bằng **khóa công khai**"], None),
    ], gap=0.20, head_size=11.6, body_size=10.4)
    yb = callout(s, x2, yb + 0.16, w2,
                 "Cloud xác định được **tài liệu nào phù hợp truy vấn** mà không cần biết "
                 "nội dung đầy đủ.", tone="green", size=10.9)
    y = max(ya, yb)
    y = callout(s, ML, y + 0.10, CW,
                "Kết phần 3 — **ABE → kiểm soát truy cập** · **FHE → tính toán trên dữ liệu "
                "mã hóa** · **SE → tìm kiếm trên dữ liệu mã hóa**.",
                tone="navy", size=11.4)
    check("Slide 10", y)


# =============================================================== SLIDE 11
def slide11(d):
    s = frame(d)
    wl = 3.20
    y = TOP
    y = label(s, ML, y, wl, "1 · Zoom out: từ dữ liệu đến hạ tầng")
    ya = layers(s, ML, y, wl, ["DATA", "NETWORK", "IDENTITY & ACCESS", "HOST / VM"],
                bh=0.50, gap=0.13, size=11.2)
    ya = callout(s, ML, ya + 0.16, wl,
                 "Bảo mật Cloud không chỉ bảo vệ dữ liệu mà còn bảo vệ **môi trường lưu trữ "
                 "và xử lý** dữ liệu.", size=10.8)

    x2 = ML + wl + 0.32
    w2 = CW - wl - 0.32
    yb = TOP
    yb = label(s, x2, yb, w2, "2 · Private Cloud vs. Public Cloud")
    yb = grid(s, x2, yb, w2, ["Private Cloud", "Public Cloud"],
              [["Mạng được tổ chức & kiểm soát tập trung hơn",
                "Tài nguyên cung cấp qua hạ tầng của CSP"],
               ["Dùng được các cơ chế kiểm soát mạng nội bộ quen thuộc",
                "Phụ thuộc nhiều hơn vào cấu hình & dịch vụ của CSP"],
               ["Vẫn phải đối phó cấu hình sai, truy cập trái phép, insider threat",
                "Tăng phụ thuộc vào kết nối mạng & cấu hình tài nguyên"]],
              first_col_bold=False, size=10.3)
    yb += 0.16
    yb = cards_row(s, x2, yb, w2, 1.02, [
        ("Network Exposure", ["Tài nguyên có thể truy cập qua Internet"], None),
        ("Misconfiguration", ["Cấu hình sai mở rộng phạm vi truy cập"], None),
        ("Visibility", ["Quan sát phụ thuộc công cụ của CSP"], None),
        ("Availability", ["Phụ thuộc kết nối & hạ tầng bên ngoài"], None),
    ], gap=0.16, head_size=10.8, body_size=10.0)
    yb = callout(s, x2, yb + 0.16, w2,
                 "**Public Cloud ≠ không an toàn** — vấn đề nằm ở cách **thiết kế, cấu hình "
                 "và kiểm soát**.", tone="amber", size=10.8)

    y = max(ya, yb) + 0.20
    y = label(s, ML, y, CW, "3 · Trust Zones: giới hạn phạm vi của một sự cố")
    y = flow_h(s, ML, y, CW, 0.48,
               ["Users", "Zone A", "Zone B", "Zone C", "Critical Resources"],
               gap=0.26, size=11.0, last=(REDBG, REDBD, RED))
    y = callout(s, ML, y + 0.10, CW,
                "**Trust Zone = Network Segmentation + IAM.** Mỗi vùng có quyền truy cập, "
                "chính sách mạng và giám sát riêng → chặn chuỗi **Initial Access → Lateral "
                "Movement → Critical Resources**. Không mặc nhiên tin tưởng chỉ vì một thành "
                "phần đang nằm bên trong mạng.", tone="navy", size=11.0)
    check("Slide 11", y)


# =============================================================== SLIDE 12
def slide12(d):
    s = frame(d)
    y = TOP
    y = label(s, ML, y, CW, "1 · Cùng là Cloud — trách nhiệm không giống nhau")
    y = grid(s, ML, y, CW, ["", "IaaS", "PaaS", "SaaS"],
             [["CSP chịu trách nhiệm", "Data center, phần cứng, mạng, hypervisor",
               "IaaS + OS / runtime / middleware của nền tảng",
               "Hạ tầng + nền tảng + ứng dụng"],
              ["Khách hàng chịu trách nhiệm", "OS, cấu hình mạng, ứng dụng, dữ liệu, IAM",
               "Mã nguồn, dữ liệu, IAM & cấu hình ứng dụng",
               "Dữ liệu, người dùng, IAM & cấu hình dịch vụ"],
              ["Mức kiểm soát & điểm cần chú ý",
               "Kiểm soát cao — VM & OS Security",
               "Kiểm soát trung bình — Application Security",
               "Kiểm soát thấp hơn — Identity & Access Control"]],
             ratios=[1.0, 1.35, 1.35, 1.30], size=9.9)
    y = callout(s, ML, y + 0.10, CW,
                "Dịch vụ ở **tầng càng cao**, CSP đảm nhận càng nhiều — nhưng khách hàng "
                "**vẫn phải quản lý phần trách nhiệm của mình**.", size=11.0)
    y += 0.22

    w2 = (CW - 0.34) / 2
    ya = label(s, ML, y, w2, "2 · Từ Network-centric → Identity-centric")
    ya = cards_row(s, ML, ya, w2, 1.06, [
        ("Truyền thống", ["**“Bạn kết nối từ đâu?”**",
                          "IP · firewall · network perimeter"], None),
        ("Trên Cloud", ["**“Bạn là ai, được phép làm gì?”**",
                        "Identity · authentication · authorization"], None),
    ], gap=0.20, head_size=11.4, body_size=10.2)
    ya = callout(s, ML, ya + 0.14, w2,
                 "IAM trở thành **lớp kiểm soát trung tâm** khi người dùng, thiết bị và "
                 "dịch vụ truy cập từ nhiều vị trí.", size=10.6)

    x2 = ML + w2 + 0.34
    yb = label(s, x2, y, w2, "3 · Quyền truy cập phải đủ fine-grained")
    yb = flow_h(s, x2, yb, w2, 0.58, [
        "Ai?|Identity", "Được làm gì?|Permission",
        "Tài nguyên nào?|Resource", "Điều kiện nào?|Context"],
        gap=0.14, size=10.2)
    yb = callout(s, x2, yb + 0.14, w2,
                 "**Mục tiêu:** đúng quyền — đúng tài nguyên — đúng thời điểm. Thu hồi "
                 "quyền một tài nguyên **không đồng nghĩa** mọi tài nguyên liên quan "
                 "cũng bị thu hồi.", tone="red", size=10.8)
    y = max(ya, yb)
    y = callout(s, ML, y + 0.10, CW,
                "**IaaS → nhiều quyền kiểm soát hơn → nhiều trách nhiệm hơn:** cấu hình → "
                "gia cố → bảo vệ khóa → firewall → logging → audit.",
                tone="navy", size=10.8)
    check("Slide 12", y)


# =============================================================== SLIDE 13
def slide13(d):
    s = frame(d)
    w2 = (CW - 0.34) / 2
    y = TOP
    y = label(s, ML, y, w2, "1 · PaaS: ai chịu trách nhiệm bảo mật?")
    y = grid(s, ML, y, w2, ["CSP quản lý", "Khách hàng quản lý"],
             [["Hạ tầng & nền tảng dịch vụ", "Mã nguồn ứng dụng"],
              ["Runtime thuộc phạm vi dịch vụ", "Dữ liệu & cấu hình ứng dụng"],
              ["Vá lỗi / bảo vệ nền tảng", "Dependency & thư viện"],
              ["Tính sẵn sàng theo SLA", "Quyền truy cập & tài khoản ứng dụng"]],
             first_col_bold=False, size=10.1)
    y = callout(s, ML, y + 0.14, w2,
                "Ranh giới trách nhiệm **phụ thuộc dịch vụ PaaS cụ thể** và cách triển khai.",
                tone="amber", size=10.8)
    y += 0.20
    y = label(s, ML, y, w2, "3 · Kiểm soát đúng ranh giới")
    y = layers(s, ML, y, w2, [
        "Infrastructure → Platform → Runtime|CSP quản lý",
        "Application|Khách hàng quản lý",
        "Data & Identity|Khách hàng quản lý",
        "3rd-party Dependencies|Cần kiểm kê & đánh giá",
    ], bh=0.40, gap=0.09, size=10.5)

    x2 = ML + w2 + 0.34
    yb = TOP
    yb = label(s, x2, yb, w2, "2 · Ba vùng cần kiểm soát")
    yb = card(s, x2, yb, w2, 1.16, "CSP — Platform Security",
              ["Bảo vệ **hạ tầng, nền tảng và runtime** thuộc phạm vi dịch vụ.",
               "Quản lý bản vá và cô lập workload theo kiến trúc dịch vụ."],
              accent=BLUE, head_size=11.8, body_size=10.2, bullet=True)
    yb = card(s, x2, yb + 0.16, w2, 1.16, "Khách hàng — Application Security",
              ["Bảo vệ **mã nguồn, dữ liệu và cấu hình ứng dụng**.",
               "Kiểm soát dependency, secret và quyền truy cập."],
              accent=BLUE, head_size=11.8, body_size=10.2, bullet=True)
    yb = card(s, x2, yb + 0.16, w2, 1.44, "Third-party Dependencies",
              ["Ứng dụng phụ thuộc **API, thư viện, dịch vụ bên ngoài**.",
               "Một thành phần yếu → **điểm xâm nhập vào chuỗi ứng dụng**.",
               "Cần kiểm kê, đánh giá và cập nhật định kỳ."],
              accent=AMB, fill=AMBBG, border=AMBBD, head_color=AMB, hl=AMB,
              head_size=11.8, body_size=10.2, bullet=True)
    y = max(y, yb)
    y = callout(s, ML, y + 0.10, CW,
                "**PaaS giảm gánh nặng hạ tầng, nhưng không làm ứng dụng tự động an "
                "toàn:** CSP bảo vệ nền tảng → khách hàng bảo vệ ứng dụng & dữ liệu → "
                "dependency phải được kiểm soát.", tone="navy", size=10.8)
    check("Slide 13", y)


# =============================================================== SLIDE 14
def slide14(d):
    s = frame(d)
    y = TOP
    y = label(s, ML, y, CW, "1 · Từ “vị trí mạng” → “định danh”")
    y = flow_h(s, ML, y, CW, 0.56, [
        "Identity|Bạn là ai?", "Authentication|Xác thực",
        "Authorization|Được phép làm gì?", "Resource|Trên tài nguyên nào?",
        "Logging|Có được ghi nhận?"],
        gap=0.22, size=10.3)
    y = callout(s, ML, y + 0.10, CW,
                "IP, port, firewall vẫn quan trọng nhưng **không đủ để xác định quyền "
                "truy cập** — **xác thực thành công ≠ được phép làm mọi thứ**.", size=10.8)
    y += 0.22

    w2 = (CW - 0.34) / 2
    ya = label(s, ML, y, w2, "2 · IAM — nguyên tắc cốt lõi")
    ya = chips(s, ML, ya, w2,
               ["MFA", "Least Privilege", "Role-based Access", "Logging"],
               size=10.2, per_row=4, gap=0.12)
    ya += 0.20
    ya = label(s, ML, ya, w2, "3 · SSO — xác thực tập trung")
    ya = flow_v(s, ML + 0.45, ya, w2 - 0.90, [
        "User", "Identity Provider (IdP)", "Authentication + MFA",
        "Token / Session", "Ứng dụng liên kết"],
        bh=0.30, gap=0.15, size=10.0, last=(GRNBG, GRNBD, GREEN))

    x2 = ML + w2 + 0.34
    yb = label(s, x2, y, w2, "Lợi ích & rủi ro của SSO")
    yb = card(s, x2, yb, w2, 1.34, "Lợi ích",
              ["Giảm số mật khẩu phải quản lý.",
               "Tập trung quản lý vòng đời tài khoản.",
               "Thu hồi quyền nhanh hơn; chính sách xác thực & giám sát tập trung."],
              fill=GRNBG, border=GRNBD, head_color=GREEN, hl=GREEN, accent=GREEN,
              head_size=12.0, body_size=10.1, bullet=True)
    yb = card(s, x2, yb + 0.14, w2, 1.52, "Rủi ro cần kiểm soát",
              ["**Compromised identity:** tài khoản trung tâm bị chiếm → nhiều dịch vụ bị "
               "ảnh hưởng.",
               "**Stolen session / token:** phiên hợp lệ bị đánh cắp.",
               "**Over-privilege:** SSO không tự bảo đảm quyền tối thiểu."],
              fill=REDBG, border=REDBD, head_color=RED, hl=RED, accent=RED,
              head_size=12.0, body_size=10.1, bullet=True)
    y = max(ya, yb)
    y = callout(s, ML, y + 0.10, CW,
                "IAM không chỉ trả lời **“ai đăng nhập?”** mà còn: **được làm gì, trên tài "
                "nguyên nào, và hoạt động đó có được ghi nhận hay không?**",
                tone="navy", size=11.2)
    check("Slide 14", y)


# =============================================================== SLIDE 15
def slide15(d):
    s = frame(d)
    w2 = (CW - 0.34) / 2
    y = TOP
    y = label(s, ML, y, w2, "1 · SaaS thay đổi trách nhiệm thế nào?")
    y = grid(s, ML, y, w2, ["SaaS Provider", "Khách hàng"],
             [["Ứng dụng & nền tảng", "Người dùng / định danh"],
              ["Hạ tầng & database", "Quyền truy cập"],
              ["Vá lỗi & vận hành dịch vụ", "Dữ liệu & cấu hình"],
              ["Bảo mật dịch vụ", "Kiểm soát việc sử dụng"]],
             first_col_bold=False, size=10.3)
    y = callout(s, ML, y + 0.14, w2,
                "CSP quản lý nhiều hơn → khách hàng ít kiểm soát hạ tầng hơn → **IAM, dữ "
                "liệu và quyền truy cập** trở nên đặc biệt quan trọng.", size=10.8)
    y += 0.20
    y = label(s, ML, y, w2, "Khi đánh giá một dịch vụ SaaS")
    y = chips(s, ML, y, w2, ["Architecture", "Secure Development", "Security Testing",
                             "Release Management", "Access Control", "Audit & Review"],
              size=10.2, per_row=2, gap=0.12)

    x2 = ML + w2 + 0.34
    yb = TOP
    yb = label(s, x2, yb, w2, "2 · Fine-grained Access Control")
    yb = card(s, x2, yb, w2, 0.86, None,
              ["Không chỉ hỏi “ai được vào ứng dụng?” mà **“ai được làm gì, trên tài nguyên "
               "nào, trong điều kiện nào?”**"], body_size=10.9, pad=0.14)
    yb = flow_h(s, x2, yb + 0.16, w2, 0.78, [
        "Identity|Nhân viên kế toán", "Resource|Báo cáo tài chính",
        "Action|Xem / sửa / chia sẻ", "Condition|Thiết bị, vị trí, thời gian"],
        gap=0.14, size=10.0)
    yb += 0.18
    yb = card(s, x2, yb, w2, 1.38, "3 · Rủi ro khi quyền quá thô",
              ["Chỉ có mức quyền đơn giản **View / Edit / Share**.",
               "Khó giới hạn theo **người dùng → tài nguyên → hành động → điều kiện**.",
               "Dễ phát sinh **over-privilege**, khó đáp ứng kiểm toán."],
              fill=REDBG, border=REDBD, head_color=RED, hl=RED, accent=RED,
              head_size=12.0, body_size=10.4, bullet=True)
    yb = callout(s, x2, yb + 0.16, w2,
                 "Tài liệu có thể **nhúng / tham chiếu tài nguyên khác**: thu hồi quyền tài "
                 "liệu chính ≠ thu hồi mọi quyền liên quan.", tone="amber", size=10.8)
    y = max(y, yb)
    y = callout(s, ML, y + 0.10, CW,
                "**Identity → Resource → Action → Condition**, đi kèm **Least Privilege + "
                "Logging + Monitoring + Access Review**.", tone="navy", size=11.4)
    check("Slide 15", y)


# =============================================================== SLIDE 16
def slide16(d):
    s = frame(d)
    y = TOP
    y = label(s, ML, y, CW, "1 · IaaS: khách hàng trực tiếp quản lý VM")
    y = flow_h(s, ML, y, CW, 0.42,
               ["VM", "OS", "Services", "Accounts", "Network Rules", "Application"],
               gap=0.22, size=10.8)
    y = callout(s, ML, y + 0.10, CW,
                "**Self-provisioning** tạo VM rất nhanh nhưng dễ sinh ra: image chưa gia "
                "cố · cổng/dịch vụ thừa · credential lưu sai chỗ · quyền quá rộng. "
                "→ **Secure-by-default: VM phải bắt đầu từ cấu hình an toàn.**",
                tone="amber", size=10.8)
    y += 0.20
    y = label(s, ML, y, CW, "2 · VM Security Checklist — 6 lớp kiểm soát")
    h = 1.00
    cards_row(s, ML, y, CW, h, [
        ("1 · Hardened Image",
         ["Golden image đã kiểm tra; chỉ cài thành phần cần thiết"], None),
        ("2 · Patch & Configuration",
         ["Vá OS/package; chuẩn hóa cấu hình, kiểm tra drift"], None),
        ("3 · Identity & Credentials",
         ["Không lưu secret trong image; MFA/SSH key; least privilege"], None),
    ], gap=0.22, head_size=11.2, body_size=10.0)
    y = cards_row(s, ML, y + h + 0.16, CW, h, [
        ("4 · Host Firewall",
         ["Chỉ mở port cần thiết; giới hạn nguồn; tắt dịch vụ thừa"], None),
        ("5 · Logging & Auditing",
         ["Thu thập sự kiện xác thực, đặc quyền, hệ thống về nơi tập trung"], None),
        ("6 · Continuous Verification",
         ["Quét lỗ hổng, đánh giá cấu hình, kiểm tra định kỳ"], None),
    ], gap=0.22, head_size=11.2, body_size=10.0)
    y += 0.20
    y = label(s, ML, y, CW, "3 · Secure VM Lifecycle — một quy trình liên tục")
    y = flow_h(s, ML, y, CW, 0.46, [
        "Golden Image", "Secure Config", "Controlled Provisioning", "Least Privilege",
        "Minimal Attack Surface", "Central Logging", "Scan · Patch · Review"],
        gap=0.12, size=9.4, last=(GRNBG, GRNBD, GREEN))
    y = callout(s, ML, y + 0.10, CW,
                "**VM Security ≠ hardening một lần** — bảo mật VM bắt đầu **trước khi VM "
                "được tạo** và tiếp tục suốt vòng đời của nó.", tone="navy", size=10.8)
    check("Slide 16", y)


# =============================================================== SLIDE 17
def slide17(d):
    s = frame(d)
    y = TOP
    y = label(s, ML, y, CW, "1 · Bốn nhóm Security Controls")
    y = cards_row(s, ML, y, CW, 1.66, [
        ("Deterrent", ["**Mục tiêu:** giảm ý định vi phạm",
                       "Security policy · warning · legal notice"], None),
        ("Preventive", ["**Mục tiêu:** ngăn / giảm khả năng xảy ra",
                        "MFA · IAM · mã hóa · firewall · hardening"], None),
        ("Detective", ["**Mục tiêu:** phát hiện bất thường",
                       "Logging · monitoring · IDS/IPS · alerting"], None,
         dict(fill=AMBBG, border=AMBBD, head_color=AMB, hl=AMB)),
        ("Corrective", ["**Mục tiêu:** giảm tác động & khôi phục",
                        "Incident response · isolation · backup · recovery"], None,
         dict(fill=GRNBG, border=GRNBD, head_color=GREEN, hl=GREEN)),
    ], gap=0.22, head_size=13.0, body_size=10.5)
    y += 0.22
    y = label(s, ML, y, CW, "2 · Các control phối hợp theo vòng đời sự cố")
    y = flow_h(s, ML, y, CW, 0.48, [
        "Security Event", "Deterrent", "Preventive", "Detective", "Corrective",
        "Recovery"], gap=0.20, size=10.6, last=(GRNBG, GRNBD, GREEN))
    y += 0.22
    y = label(s, ML, y, CW, "3 · Ví dụ: tài khoản Cloud bị chiếm quyền")
    y = cards_row(s, ML, y, CW, 0.94, [
        ("Preventive", ["MFA + least privilege"], None),
        ("Detective", ["Phát hiện đăng nhập bất thường + logging"], None,
         dict(fill=AMBBG, border=AMBBD, head_color=AMB, hl=AMB)),
        ("Corrective", ["Thu hồi session/token → khóa tài khoản → reset credential"], None,
         dict(fill=GRNBG, border=GRNBD, head_color=GREEN, hl=GREEN)),
        ("Recovery", ["Khôi phục cấu hình / dữ liệu nếu bị thay đổi"], None,
         dict(fill=GRNBG, border=GRNBD, head_color=GREEN, hl=GREEN)),
    ], gap=0.22, head_size=11.6, body_size=10.3)
    y = callout(s, ML, y + 0.16, CW,
                "Mục tiêu không chỉ là “ngăn chặn” mà là một **chuỗi kiểm soát liên tục: "
                "Deterrent → Preventive → Detective → Corrective** — không control đơn lẻ "
                "nào bảo vệ được toàn bộ hệ thống.", tone="navy", size=11.2)
    check("Slide 17", y)


# =============================================================== SLIDE 18
def slide18(d):
    s = frame(d)
    wl = 4.20
    y = TOP
    y = label(s, ML, y, wl, "1 · Kiến trúc bảo mật nhiều lớp")
    yl = layers(s, ML, y, wl, [
        "SHARED RESPONSIBILITY + CIA|Nền tảng",
        "DATA SECURITY|Crypto + access control",
        "CRYPTOGRAPHY · IAM · NETWORK|ABE/FHE/SE · identity · trust zones",
        "APPLICATION & VM|SaaS · PaaS · IaaS",
        "MONITORING & LOGGING|Phát hiện",
        "RESPONSE & RECOVERY|Khắc phục",
        "DEFENSE-IN-DEPTH|Kết quả tổng hợp",
    ], bh=0.58, gap=0.11, size=10.6, top_tone=(NAVY, NAVY, WHITE))

    x2 = ML + wl + 0.34
    w2 = CW - wl - 0.34
    y = TOP
    y = label(s, x2, y, w2, "2 · Ba nguyên tắc cần nhớ")
    y = card(s, x2, y, w2, 1.22, "1 · Shared Responsibility",
             ["Cloud **không loại bỏ** trách nhiệm bảo mật — Cloud **phân chia lại**. "
              "**IaaS:** OS, VM, mạng, app, dữ liệu · **PaaS:** app, dữ liệu, identity · "
              "**SaaS:** identity, dữ liệu, quyền truy cập."],
             accent=BLUE, head_size=12.2, body_size=10.4)
    y = card(s, x2, y + 0.14, w2, 1.10, "2 · Data Security ≠ Encryption only",
             ["Kết hợp **CIA** với **access control + privacy + data lifecycle + "
              "metadata**: mã hóa bảo vệ dữ liệu, nhưng **quyền truy cập** quyết định ai "
              "dùng được dữ liệu."],
             accent=BLUE, head_size=12.2, body_size=10.4)
    y = card(s, x2, y + 0.14, w2, 1.10, "3 · Defense-in-Depth",
             ["Không phụ thuộc một cơ chế duy nhất: **crypto → IAM → network → "
              "application/VM → monitoring → response**. Một lớp bị vượt qua, các lớp còn "
              "lại vẫn giới hạn tác động."],
             accent=BLUE, head_size=12.2, body_size=10.4)
    y = card(s, x2, y + 0.14, w2, 1.34, "Đặc biệt bảo vệ: Credentials & Secrets",
             ["API keys · passwords · tokens · certificates · SSH keys · service accounts.",
              "**Leaked secret → unauthorized access → privilege escalation → lateral "
              "movement.**",
              "**Least Privilege → Secure Storage → Rotation → Monitoring → Revocation.**"],
             fill=REDBG, border=REDBD, head_color=RED, hl=RED, accent=RED,
             head_size=12.2, body_size=10.2)
    check("Slide 18", max(y, yl))


# ===================================================================== main
BUILDERS = {1: slide1, 2: slide2, 3: slide3, 4: slide4, 5: slide5, 6: slide6,
            7: slide7, 8: slide8, 9: slide9, 10: slide10, 11: slide11,
            12: slide12, 13: slide13, 14: slide14, 15: slide15, 16: slide16,
            17: slide17, 18: slide18}


def main():
    data = {d["num"]: d for d in parse()}
    for num in range(1, 19):
        d = data[num]
        BUILDERS[num](d)
        print(f"  Slide {num:>2}  · notes {len(d['notes']):>5} ký tự  · "
              f"{TITLES.get(num, d['title'])[:52]}")
    prs.save(OUT)
    print(f"\nĐã ghi: {OUT.relative_to(ROOT)} ({len(prs.slides._sldIdLst)} slide)")
    if warnings:
        print("\nCảnh báo bố cục:")
        for w in warnings:
            print("  - " + w)
    else:
        print("Không có cảnh báo bố cục.")


if __name__ == "__main__":
    main()
