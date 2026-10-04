#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Bản RÚT GỌN 12 slide cho buổi trình bày 20-30 phút.

Nguyên tắc:
  * Mỗi slide chỉ giữ ý chính, không ví dụ, không đi sâu kỹ thuật.
  * Speaker notes viết ngắn (~1,5-2 phút/slide), không lấy nguyên văn
    presentation.md.
  * Dùng lại toàn bộ helper vẽ của tools/build_presentation.py.

Chạy:  .venv/bin/python tools/build_short.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_presentation as bp  # noqa: E402
from build_presentation import (  # noqa: E402
    BLUE, BLUEL, GREEN, GRNBD, GRNBG, MUTED, NAVY2, PALE,
    SKY2, WHITE, AMB, RED, REDBD, REDBG,
    ML, MR, CW, TOP, SW, SH,
    callout, card, cards_row, chips, flow_h, flow_v, grid, label, layers,
    new_slide, rect, rich, steps, txbox,
)
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from short_notes import NOTES  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Chapter11_Vietnamese_Presentation_Short.pptx"
TOTAL = 12


def footer(slide, idx):
    tf = txbox(slide, ML, 7.13, 8.0, 0.24)
    rich(tf.paragraphs[0], "Bảo mật đám mây (Cloud Security) · Nông Văn Tình",
         9, MUTED, line_spacing=1.0)
    tf = txbox(slide, SW - MR - 1.6, 7.13, 1.6, 0.24)
    rich(tf.paragraphs[0], f"{idx}/{TOTAL}", 9, MUTED, align=PP_ALIGN.RIGHT,
         line_spacing=1.0)


def frame(idx, kicker, title, note):
    s = new_slide()
    bp.head(s, kicker, title)
    footer(s, idx)
    if note:
        s.notes_slide.notes_text_frame.text = note.strip()
    return s


# ================================================================ SLIDE 1


def slide1():
    s = new_slide()
    rect(s, 0, 0, SW, SH, fill=bp.NAVY, shape=MSO_SHAPE.RECTANGLE)
    rect(s, SW - 4.3, -1.7, 5.6, 5.6, fill=NAVY2, shape=MSO_SHAPE.OVAL)
    rect(s, SW - 2.6, 4.2, 4.4, 4.4, fill=NAVY2, shape=MSO_SHAPE.OVAL)
    rect(s, 0, 0, 0.14, SH, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    tf = txbox(s, 1.15, 1.95, 9.0, 0.4)
    rich(tf.paragraphs[0], "CHƯƠNG 11", 17, BLUEL, bold=True, line_spacing=1.0)
    tf = txbox(s, 1.15, 2.43, 10.4, 0.9)
    rich(tf.paragraphs[0], "BẢO MẬT ĐÁM MÂY", 48, WHITE, bold=True,
         hl=WHITE, line_spacing=1.0)
    tf = txbox(s, 1.15, 3.53, 10.0, 0.45)
    rich(tf.paragraphs[0], "Cloud Security", 20, BLUEL, italic=True,
         line_spacing=1.0)
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.15),
                                Inches(4.23), Inches(3.4), Inches(4.23))
    ln.line.color.rgb = BLUE
    ln.line.width = Pt(3)
    tf = txbox(s, 1.15, 4.51, 10.2, 0.4)
    rich(tf.paragraphs[0],
         "Trách nhiệm · Dữ liệu · Mật mã · Hạ tầng · Ứng dụng · Kiểm soát",
         13.5, PALE, line_spacing=1.2)
    tf = txbox(s, 1.15, 6.18, 8.0, 0.35)
    rich(tf.paragraphs[0], "Người thực hiện: **Nông Văn Tình**", 13, PALE,
         hl=WHITE, line_spacing=1.1)
    tf = txbox(s, SW - MR - 1.6, 7.13, 1.6, 0.24)
    rich(tf.paragraphs[0], f"1/{TOTAL}", 9, RGBColor.from_string("6E8FAE"),
         align=PP_ALIGN.RIGHT, line_spacing=1.0)
    s.notes_slide.notes_text_frame.text = NOTES[1].strip()


# ================================================================ SLIDE 2


def slide2():
    s = frame(2, "NỘI DUNG TRÌNH BÀY", "MỤC LỤC", NOTES[2])
    specs = [
        ("1 · NỀN TẢNG BẢO MẬT ĐÁM MÂY",
         ["Shared Responsibility Model", "Bộ ba bảo mật CIA"], BLUE),
        ("2 · BẢO VỆ DỮ LIỆU TRÊN CLOUD",
         ["Ba trạng thái của dữ liệu", "Bốn vấn đề cần kiểm soát"], BLUE),
        ("3 · MẬT MÃ NÂNG CAO CHO CLOUD",
         ["ABE · FHE · Searchable Encryption", "Ba bài toán khác nhau"], BLUE),
        ("4 · HẠ TẦNG & VÙNG TIN CẬY",
         ["Public Cloud vs Private Cloud", "Phân đoạn mạng & Trust Zone"], BLUEL),
        ("5 · ỨNG DỤNG, DANH TÍNH & MÁY CHỦ",
         ["IaaS · PaaS · SaaS", "IAM · SSO · Máy chủ ảo"], BLUEL),
        ("6 · KIỂM SOÁT & TỔNG KẾT",
         ["Bốn nhóm kiểm soát bảo mật", "Defense in Depth"], BLUEL),
    ]
    h = 1.30
    y = cards_row(s, ML, TOP + 0.34, CW, h, specs[:3], gap=0.26, bullet=True,
                  head_size=12.5, body_size=11.2)
    y = cards_row(s, ML, y + 0.30, CW, h, specs[3:], gap=0.26, bullet=True,
                  head_size=12.5, body_size=11.2)
    y = callout(s, ML, y + 0.38, CW,
                "Mạch bài: **trách nhiệm & mục tiêu** → **dữ liệu** → **mật mã nâng cao** "
                "→ **hạ tầng, ứng dụng, máy chủ** → **các lớp kiểm soát**.")
    bp.check("Slide 2", y)


# ================================================================ SLIDE 3


def slide3():
    s = frame(3, "PHẦN 1 · NỀN TẢNG", "TRÁCH NHIỆM & MỤC TIÊU BẢO MẬT", NOTES[3])
    y = TOP + 0.20
    y = label(s, ML, y, CW,
              "Shared Responsibility Model — ranh giới dịch chuyển theo mô hình dịch vụ")
    y = cards_row(s, ML, y + 0.06, CW, 1.28, [
        ("NHÀ CUNG CẤP (CSP)",
         ["Hạ tầng vật lý, ảo hóa, mạng lõi", "Dịch vụ nền tảng họ vận hành"],
         BLUE, {"bullet": True}),
        ("KHÁCH HÀNG",
         ["Dữ liệu, danh tính, quyền truy cập", "Cấu hình dịch vụ và ứng dụng"],
         BLUEL, {"bullet": True}),
        ("DỊCH CHUYỂN",
         ["IaaS → PaaS → SaaS: CSP lo nhiều hơn",
          "Nhưng khách hàng không hết trách nhiệm"],
         AMB, {"bullet": True}),
    ], gap=0.26, head_size=12, body_size=11)

    y = label(s, ML, y + 0.34, CW, "Mục tiêu bảo mật — bộ ba CIA")
    y = cards_row(s, ML, y + 0.06, CW, 0.98, [
        ("Confidentiality", ["Ai được phép **xem** dữ liệu?"], None),
        ("Integrity", ["Dữ liệu có bị **thay đổi** trái phép?"], None),
        ("Availability", ["Khi cần, hệ thống có **sẵn sàng**?"], None),
    ], gap=0.26, head_size=12.5, body_size=11.5)

    y = callout(s, ML, y + 0.30, CW,
                "Cloud **không xóa bỏ** trách nhiệm bảo mật — nó **phân chia lại** trách nhiệm.")
    bp.check("Slide 3", y)


# ================================================================ SLIDE 4


def slide4():
    s = frame(4, "PHẦN 2 · BẢO VỆ DỮ LIỆU",
              "DỮ LIỆU TRÊN CLOUD ĐƯỢC BẢO VỆ THẾ NÀO?", NOTES[4])
    y = TOP + 0.20
    y = flow_h(s, ML, y, CW, 0.56,
               ["Người dùng", "Đường truyền", "Lưu trữ Cloud", "Xử lý Cloud",
                "Trả kết quả"], gap=0.30, size=11.5)

    y = cards_row(s, ML, y + 0.32, CW, 1.15, [
        ("Data-in-Transit", ["Đang truyền trên mạng",
                             "→ Mã hóa + xác thực đầu cuối"], BLUE, {"bullet": True}),
        ("Data-at-Rest", ["Đang nằm trong kho lưu trữ",
                          "→ Mã hóa lưu trữ + quản lý khóa"], BLUE, {"bullet": True}),
        ("Data-in-Use", ["Đang được xử lý", "→ Bài toán khó nhất"],
         RED, {"bullet": True}),
    ], gap=0.26, head_size=12, body_size=11)

    y = label(s, ML, y + 0.32, CW, "Bốn vấn đề khi dữ liệu nằm trong tay nhà cung cấp")
    y = chips(s, ML, y + 0.06, CW,
              ["Integrity — toàn vẹn", "Privacy — riêng tư",
               "Data location — lưu ở đâu", "Availability — sẵn sàng"],
              size=11.5, h=0.42)

    y = callout(s, ML, y + 0.28, CW,
                "Không chỉ nội dung cần được bảo vệ: **metadata** cũng tiết lộ thông tin. "
                "Và mã hóa thông thường vẫn phải **giải mã khi xử lý**.", tone="amber")
    bp.check("Slide 4", y)


# ================================================================ SLIDE 5


def slide5():
    s = frame(5, "PHẦN 3 · MẬT MÃ NÂNG CAO", "BA BÀI TOÁN — BA KỸ THUẬT", NOTES[5])
    y = TOP + 0.24
    y = flow_h(s, ML, y, CW, 0.60,
               ["Mã hóa|phía người dùng", "Tải lên|Cloud", "Giải mã|để xử lý",
                "Cloud thấy|dữ liệu gốc"],
               gap=0.34, size=11.5, last=(REDBG, REDBD, RED))

    y = callout(s, ML, y + 0.26, CW,
                "Bài toán: **muốn Cloud xử lý dữ liệu, nhưng không muốn Cloud biết dữ liệu.**",
                tone="red", size=12.5)

    y = grid(s, ML, y + 0.34, CW,
             ["CÂU HỎI ĐẶT RA", "KỸ THUẬT", "Ý TƯỞNG"],
             [["Ai được phép giải mã?", "ABE — Attribute-Based Encryption",
               "Quyền đọc gắn với **thuộc tính** và **chính sách**"],
              ["Tính toán trên bản mã được không?", "FHE — Fully Homomorphic Encryption",
               "Cloud tính **trực tiếp trên dữ liệu đã mã hóa**"],
              ["Tìm kiếm mà không giải mã?", "SE — Searchable Encryption",
               "Tìm qua **chỉ mục mã hóa** và **trapdoor**"]],
             ratios=[1.0, 1.15, 1.35], size=11.5, head_size=11.5)
    bp.check("Slide 5", y)


# ================================================================ SLIDE 6


def slide6():
    s = frame(6, "PHẦN 3 · MẬT MÃ NÂNG CAO",
              "ABE — MÃ HÓA DỰA TRÊN THUỘC TÍNH", NOTES[6])
    y = TOP + 0.22
    y = callout(s, ML, y, CW,
                "Bài toán: **quyền đọc gắn với ai?** — Thay vì mã hóa riêng cho từng người, "
                "dữ liệu được mã hóa kèm **một chính sách truy cập**.")

    y = steps(s, ML, y + 0.32, CW, 1.05, [
        "Thuộc tính người dùng|Vai trò, phòng ban, tổ chức… được cấp trong khóa",
        "Chính sách trên dữ liệu|Điều kiện phải thỏa mới được đọc",
        "Khớp thuộc tính|Thỏa chính sách → giải mã thành công",
    ], gap=0.24, size=11.5)

    y = cards_row(s, ML, y + 0.32, CW, 1.05, [
        ("CP-ABE", ["Chính sách nằm ở **bản mã**",
                    "Người mã hóa quyết định ai được đọc"], BLUE, {"bullet": True}),
        ("KP-ABE", ["Chính sách nằm ở **khóa**",
                    "Cơ quan cấp khóa quyết định đọc được gì"], BLUEL, {"bullet": True}),
    ], gap=0.26, head_size=12.5, body_size=11.2)

    y = callout(s, ML, y + 0.28, CW,
                "Quyền truy cập nằm **ngay trong bản mã** — không phụ thuộc vào việc "
                "máy chủ có kiểm soát đúng hay không.", tone="green")
    bp.check("Slide 6", y)


# ================================================================ SLIDE 7


def slide7():
    s = frame(7, "PHẦN 3 · MẬT MÃ NÂNG CAO",
              "FHE — TÍNH TOÁN TRÊN DỮ LIỆU MÃ HÓA", NOTES[7])
    y = TOP + 0.22
    y = callout(s, ML, y, CW,
                "Bài toán: **Cloud tính toán giúp, nhưng không được nhìn thấy dữ liệu gốc.**")

    y = flow_h(s, ML, y + 0.32, CW, 0.78,
               ["Mã hóa|tại máy người dùng", "Gửi lên|Cloud",
                "Tính toán|ngay trên bản mã", "Kết quả|vẫn đang mã hóa",
                "Giải mã|chỉ người giữ khóa"],
               gap=0.26, size=11, last=(GRNBG, GRNBD, GREEN))

    y = cards_row(s, ML, y + 0.34, CW, 1.20, [
        ("Điểm đặc biệt", ["Dữ liệu gốc **không bao giờ** xuất hiện ở phía Cloud",
                           "Cloud chỉ **thực hiện phép tính**"], GREEN, {"bullet": True}),
        ("Đánh đổi", ["Chi phí tính toán **lớn hơn nhiều**",
                      "Nhiễu tích lũy → cần **bootstrapping**"], AMB, {"bullet": True}),
    ], gap=0.26, head_size=12.5, body_size=11.2)

    y = callout(s, ML, y + 0.28, CW,
                "FHE biến Cloud từ **“nơi nhìn thấy dữ liệu”** thành "
                "**“nơi thực hiện phép tính trên dữ liệu”**.", tone="green")
    bp.check("Slide 7", y)


# ================================================================ SLIDE 8


def slide8():
    s = frame(8, "PHẦN 3 · MẬT MÃ NÂNG CAO",
              "SE — TÌM KIẾM TRÊN DỮ LIỆU MÃ HÓA", NOTES[8])
    y = TOP + 0.22
    y = callout(s, ML, y, CW,
                "Bài toán: mã hóa hết thì **Cloud không tìm kiếm được** — "
                "mà tải cả kho về giải mã thì không khả thi.", tone="red")

    half = (CW - 0.34) / 2
    y1 = label(s, ML, y + 0.32, half, "Khi tải dữ liệu lên")
    y1 = flow_v(s, ML, y1 + 0.06, half,
                ["Rút **từ khóa** từ tài liệu",
                 "Tạo **chỉ mục** — cũng được mã hóa",
                 "Lưu dữ liệu mã hóa + chỉ mục lên Cloud"],
                bh=0.46, gap=0.18, size=11)

    x2 = ML + half + 0.34
    y2 = label(s, x2, y + 0.32, half, "Khi tìm kiếm")
    y2 = flow_v(s, x2, y2 + 0.06, half,
                ["Người dùng gửi **trapdoor** (truy vấn đã mã hóa)",
                 "Cloud **đối chiếu** trapdoor với chỉ mục",
                 "Trả về đúng tài liệu — **không đọc nội dung**"],
                bh=0.46, gap=0.18, size=11, last=(GRNBG, GRNBD, GREEN))

    y = max(y1, y2)
    y = callout(s, ML, y + 0.28, CW,
                "Lưu ý: vẫn có thể **rò rỉ gián tiếp** (mẫu truy cập) → đây là hướng "
                "còn đang được nghiên cứu cải tiến.", tone="amber")
    bp.check("Slide 8", y)


# ================================================================ SLIDE 9


def slide9():
    s = frame(9, "PHẦN 4 · HẠ TẦNG", "BẢO VỆ HẠ TẦNG & VÙNG TIN CẬY", NOTES[9])
    y = TOP + 0.20
    half = (CW - 0.34) / 2

    y1 = label(s, ML, y, half, "Bảo mật theo lớp")
    y1 = layers(s, ML, y1 + 0.06, half,
                ["Dữ liệu|Mã hóa · phân loại",
                 "Mạng|Phân đoạn · firewall",
                 "Danh tính & truy cập|IAM · least privilege",
                 "Máy chủ / Máy ảo|Hardening · giám sát"],
                bh=0.62, gap=0.10, size=11.5)

    x2 = ML + half + 0.34
    y2 = label(s, x2, y, half, "Bối cảnh khác nhau")
    y2 = cards_row(s, x2, y2 + 0.06, half, 1.05, [
        ("Private Cloud", ["Kiểm soát chặt hơn", "Gần mô hình mạng nội bộ"],
         BLUE, {"bullet": True}),
        ("Public Cloud", ["Kết nối Internet, dùng chung", "Phụ thuộc cấu hình đúng"],
         BLUEL, {"bullet": True}),
    ], gap=0.20, head_size=11.5, body_size=10.6)
    y2 = label(s, x2, y2 + 0.24, half, "Nguy cơ điển hình")
    y2 = chips(s, x2, y2 + 0.04, half,
               ["APT", "Insider threat", "Truy cập trái phép", "Lateral movement"],
               size=10.6, h=0.36, per_row=2)

    y = max(y1, y2)
    y = callout(s, ML, y + 0.26, CW,
                "Giải pháp: **phân đoạn mạng** + **vùng tin cậy (trust zone)** + **IAM**. "
                "Không phải cứ ở bên trong mạng thì mặc nhiên đáng tin cậy.")
    bp.check("Slide 9", y)


# =============================================================== SLIDE 10


def slide10():
    s = frame(10, "PHẦN 5 · ỨNG DỤNG & DANH TÍNH",
              "IaaS / PaaS / SaaS & QUẢN LÝ TRUY CẬP", NOTES[10])
    y = TOP + 0.20
    y = grid(s, ML, y, CW,
             ["", "IaaS", "PaaS", "SaaS"],
             [["CSP quản lý", "Hạ tầng", "Hạ tầng + nền tảng",
               "Gần như toàn bộ ứng dụng"],
              ["Khách hàng lo", "OS, mạng, ứng dụng", "Ứng dụng, dữ liệu, truy cập",
               "Người dùng, dữ liệu, phân quyền"],
              ["Rủi ro tiêu biểu", "Cấu hình máy ảo", "Nền tảng & truy cập",
               "Phân quyền & cấu hình sai"]],
             ratios=[0.85, 1.0, 1.15, 1.25], size=11, head_size=11.5)

    y = label(s, ML, y + 0.32, CW,
              "Xuyên suốt cả ba mô hình: Identity & Access Management")
    y = cards_row(s, ML, y + 0.06, CW, 1.05, [
        ("SSO", ["Đăng nhập một lần", "Tập trung việc xác thực"],
         BLUE, {"bullet": True}),
        ("Xác thực mạnh", ["Đa yếu tố (MFA)", "Chống lộ mật khẩu"],
         BLUE, {"bullet": True}),
        ("Least Privilege", ["Chỉ cấp quyền cần dùng", "Rà soát định kỳ"],
         BLUEL, {"bullet": True}),
        ("Phân quyền chi tiết", ["Theo vai trò, theo tài nguyên",
                                 "Quản lý secrets & API key"], BLUEL, {"bullet": True}),
    ], gap=0.22, head_size=11.5, body_size=10.4)

    y = callout(s, ML, y + 0.26, CW,
                "Trên Cloud, sự cố thường không đến từ việc **phá vỡ mã hóa**, "
                "mà đến từ **cấu hình sai** và **quyền cấp thừa**.", tone="amber")
    bp.check("Slide 10", y)


# =============================================================== SLIDE 11


def slide11():
    s = frame(11, "PHẦN 6 · MÁY CHỦ & KIỂM SOÁT",
              "MÁY CHỦ ẢO & CÁC LỚP KIỂM SOÁT BẢO MẬT", NOTES[11])
    y = TOP + 0.20
    y = label(s, ML, y, CW, "Máy chủ ảo — 5 việc tối thiểu")
    y = chips(s, ML, y + 0.06, CW,
              ["Cấu hình an toàn mặc định", "Chỉ dùng image đã kiểm tra",
               "Không nhúng khóa / chứng chỉ", "Bật firewall trên máy ảo",
               "Log tập trung ra ngoài"], size=10.8, h=0.44)

    y = label(s, ML, y + 0.34, CW, "Bốn nhóm kiểm soát — Defense in Depth")
    y = cards_row(s, ML, y + 0.06, CW, 1.30, [
        ("1 · Răn đe (Deterrent)", ["Chính sách, cảnh báo",
                                    "Làm nản lòng kẻ tấn công"], BLUEL, {"bullet": True}),
        ("2 · Ngăn ngừa (Preventive)", ["Mã hóa, IAM, firewall",
                                        "Hardening hệ thống"], BLUE, {"bullet": True}),
        ("3 · Phát hiện (Detective)", ["Giám sát, log",
                                       "Phát hiện xâm nhập"], AMB, {"bullet": True}),
        ("4 · Khắc phục (Corrective)", ["Sao lưu, khôi phục",
                                        "Ứng phó sự cố"], GREEN, {"bullet": True}),
    ], gap=0.22, head_size=11.5, body_size=10.6)

    y = callout(s, ML, y + 0.26, CW,
                "Không lớp nào đủ một mình — hiệu quả đến từ việc **kết hợp cả bốn nhóm**.")
    bp.check("Slide 11", y)


# =============================================================== SLIDE 12


def slide12():
    s = frame(12, "TỔNG KẾT", "CLOUD SECURITY LÀ MỘT HỆ THỐNG NHIỀU LỚP", NOTES[12])
    y = TOP + 0.24
    y = flow_h(s, ML, y, CW, 0.62,
               ["Trách nhiệm", "Dữ liệu", "Mật mã", "Mạng & Trust zone",
                "IAM & Ứng dụng", "Giám sát & Khôi phục"],
               gap=0.20, size=10.8, last=(GRNBG, GRNBD, GREEN))

    y = cards_row(s, ML, y + 0.38, CW, 1.60, [
        ("1 · Phân chia lại trách nhiệm",
         ["Cloud không loại bỏ trách nhiệm bảo mật; phần của khách hàng luôn còn đó."],
         BLUE, {}),
        ("2 · Bảo vệ dữ liệu ≠ chỉ mã hóa",
         ["Còn là quyền truy cập, toàn vẹn, vị trí lưu trữ và khả năng xử lý an toàn."],
         BLUE, {}),
        ("3 · Nhiều lớp kết hợp",
         ["Mật mã, IAM, mạng, ứng dụng, giám sát và khôi phục — Defense in Depth."],
         BLUEL, {}),
    ], gap=0.26, head_size=12, body_size=11.2)

    y = callout(s, ML, y + 0.34, CW,
                "Xin cảm ơn thầy cô và các bạn đã lắng nghe — "
                "**rất mong nhận được câu hỏi và góp ý.**", tone="navy", size=13)
    bp.check("Slide 12", y)


def main():
    for f in (slide1, slide2, slide3, slide4, slide5, slide6, slide7, slide8,
              slide9, slide10, slide11, slide12):
        f()
    bp.prs.save(OUT)
    print(f"Đã ghi: {OUT.relative_to(ROOT)} "
          f"({len(bp.prs.slides._sldIdLst)} slide)")
    if bp.warnings:
        print("Cảnh báo bố cục:")
        for w in bp.warnings:
            print("  - " + w)
    else:
        print("Không có cảnh báo bố cục.")


if __name__ == "__main__":
    main()
